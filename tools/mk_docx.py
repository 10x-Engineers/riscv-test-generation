#!/usr/bin/env python3
"""Markdown -> .docx for Google Docs.

Google Docs opens .docx natively and keeps headings, tables and lists editable,
so this is the reliable route into docs.google.com. Handles the subset the
project's documents actually use: ATX headings, paragraphs with bold/inline
code, bullet lists, pipe tables, fenced blocks and rules.
"""
import re
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor, Inches

NAVY = RGBColor(0x0F, 0x30, 0x57)
INK = RGBColor(0x2B, 0x2F, 0x36)
MUTED = RGBColor(0x5C, 0x66, 0x72)


def add_runs(par, text, size=10.5, color=INK):
    """Render **bold**, `code` and plain text as separate runs."""
    for part in re.split(r"(\*\*.+?\*\*|`[^`]+`|\*[^*]+\*)", text):
        if not part:
            continue
        r = par.add_run()
        if part.startswith("**") and part.endswith("**"):
            r.text = part[2:-2]; r.bold = True; r.font.color.rgb = NAVY
        elif part.startswith("`") and part.endswith("`"):
            r.text = part[1:-1]; r.font.name = "Consolas"; r.font.color.rgb = MUTED
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            r.text = part[1:-1]; r.italic = True; r.font.color.rgb = color
        else:
            r.text = part; r.font.color.rgb = color
        r.font.size = Pt(size)


def convert(src, out, title=None):
    lines = open(src).read().split("\n")
    doc = Document()
    for s in doc.sections:
        s.left_margin = s.right_margin = Inches(0.9)
        s.top_margin = s.bottom_margin = Inches(0.8)
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10.5)

    i, fence = 0, False
    while i < len(lines):
        l = lines[i]

        if l.startswith("```"):
            fence = not fence
            if fence:  # gather the block
                block = []
                i += 1
                while i < len(lines) and not lines[i].startswith("```"):
                    block.append(lines[i]); i += 1
                fence = False
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.2)
                p.paragraph_format.space_after = Pt(8)
                r = p.add_run("\n".join(block))
                r.font.name = "Consolas"; r.font.size = Pt(8); r.font.color.rgb = MUTED
            i += 1
            continue

        if m := re.match(r"^(#{1,4})\s+(.*)", l):
            lvl, txt = len(m.group(1)), m.group(2)
            txt = re.sub(r"[*`]", "", txt)
            h = doc.add_heading(level=min(lvl, 4))
            r = h.runs[0] if h.runs else h.add_run()
            r.text = txt
            r.font.color.rgb = NAVY
            r.font.size = Pt({1: 18, 2: 14, 3: 12, 4: 11}[min(lvl, 4)])
            i += 1
            continue

        # pipe table
        if l.startswith("|") and i + 1 < len(lines) and re.match(r"^\|[\s:|-]+\|$", lines[i + 1]):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                if not re.match(r"^\|[\s:|-]+\|$", lines[i]):
                    rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
                i += 1
            if rows:
                ncol = max(len(r) for r in rows)
                t = doc.add_table(rows=len(rows), cols=ncol)
                t.style = "Table Grid"
                t.alignment = WD_TABLE_ALIGNMENT.CENTER
                for ri, row in enumerate(rows):
                    for ci in range(ncol):
                        cell = t.cell(ri, ci)
                        cell.text = ""
                        p = cell.paragraphs[0]
                        p.paragraph_format.space_after = Pt(2)
                        add_runs(p, row[ci] if ci < len(row) else "", size=9)
                        if ri == 0:
                            for r in p.runs:
                                r.bold = True; r.font.color.rgb = NAVY
                doc.add_paragraph().paragraph_format.space_after = Pt(4)
            continue

        if re.match(r"^[-*]\s+", l):
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.space_after = Pt(3)
            add_runs(p, re.sub(r"^[-*]\s+", "", l))
            i += 1
            continue

        if re.match(r"^\d+\.\s+", l):
            p = doc.add_paragraph(style="List Number")
            p.paragraph_format.space_after = Pt(3)
            add_runs(p, re.sub(r"^\d+\.\s+", "", l))
            i += 1
            continue

        if l.strip() == "---":
            doc.add_paragraph("_" * 60).runs[0].font.color.rgb = RGBColor(0xDF, 0xE4, 0xEA)
            i += 1
            continue

        if l.strip():
            # join wrapped lines into one paragraph
            buf = [l]
            while (i + 1 < len(lines) and lines[i + 1].strip()
                   and not re.match(r"^(#{1,4}\s|\||[-*]\s|\d+\.\s|```|---$)", lines[i + 1])):
                i += 1; buf.append(lines[i])
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(7)
            add_runs(p, " ".join(x.strip() for x in buf))
        i += 1

    doc.save(out)
    return out


if __name__ == "__main__":
    for src, out in [(sys.argv[1], sys.argv[2])]:
        p = convert(src, out)
        import os
        print(f"  {os.path.basename(p)}  ({os.path.getsize(p)/1024:.0f} KB)")
