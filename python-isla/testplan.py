#!/usr/bin/env python3
"""Derive a testplan from the Sail model's own branch structure.

`coverage_report.py` answers "how much of the model do we reach" and
`coverage_analysis.py` turns the uncovered remainder into a stimulus queue.
Neither produces a *plan*: a plan has to exist before any test is run, has to
enumerate the whole target set rather than the current shortfall, and has to
name each target in terms a reviewer recognises -- an instruction, a privilege
transition, a CSR field -- rather than a file and a line number.

This builds that. The target set is the model's own instrumented span manifest
(`sail_riscv_model.branch_info`), so the plan is complete by construction: it
cannot omit a behaviour the model implements, because every behaviour the model
implements is compiled into a span. Coverage, when supplied, is joined in as a
*status column*, which is the direction that keeps the plan honest -- the plan
does not shrink when tests are added, it fills in.

The reasoning that turns a span into a plan item happens in four steps, each
derived from the sources rather than hand-maintained:

  1. Ownership.  Every span falls inside exactly one top-level Sail definition.
     `function clause execute RTYPE(...)` owns its spans; so does `pmpCheck`.
     F-kind spans in branch_info carry the function name already, which is used
     to cross-check the containment result rather than to replace it (F spans
     are 1941 of 41407; containment has to work for the other 39466 regardless).

  2. Discrimination.  Inside an execute clause the model discriminates
     instructions by a `match` on the operand enum:

         function clause execute RTYPE(rs2, rs1, rd, op) = {
           X(rd) = match op {
             ADD  => X(rs1) + X(rs2),
             SRA  => shift_bits_right_arith(...),

     Each arm is a separate span. So an uncovered span at `base_insts.sail:246`
     is not "an expression in RTYPE" -- it is *`sra` was never executed*. This
     step recovers the enclosing arm by a brace-depth scan.

  3. Naming.  The arm's pattern is an enum constructor (`SRA`), and the model
     already states the constructor's assembly text in a bidirectional mnemonic
     mapping (`SRA <-> "sra"`). Those mappings are unreachable *as code* -- see
     derive_span_exclusions.py, which removes them from the denominator for
     exactly that reason -- but they are the model's own authority on what each
     constructor is called, which makes them the right source for naming plan
     items. Excluded from measurement, used for description.

  4. Stimulus.  Where no arm applies (privileged code, memory, traps), the
     enclosing guard is read out of the source and classified into what a
     generator must arrange: a privilege level, a CSR field, a configuration
     key, an XLEN. Unrecognised guards are reported as `unclassified` rather
     than guessed at, because a wrong directive costs a generation cycle and an
     honest gap costs a minute of reading.

Plan items are then the grouping of spans by (feature, owner, target), which is
the unit a person actually schedules: "cover `sra`" is one item of four spans,
not four items.

What this deliberately does not do: claim a span is impossible. Exclusions come
from derive_span_exclusions.py, which reasons structurally about mapping
directions; anything else that is uncovered is a work item, however awkward.

Usage:
    # the plan alone, no coverage -- the complete target set
    testplan.py -o testplan

    # the plan with current status joined in
    coverage_report.py                       # refreshes sail_coverage
    derive_span_exclusions.py -o excl.txt
    testplan.py --coverage ~/Documents/sail-riscv/sail_coverage \
                --exclude-spans excl.txt --scope rfp-scope.txt -o testplan --xlsx
"""

import argparse
import collections
import csv
import json
import os
import re
import sys

import paths

try:
    from coverage_report import parse_spans
except ImportError:                                            # pragma: no cover
    sys.exit("testplan.py must run from python-isla/ (it imports coverage_report)")


# --------------------------------------------------------------------------
# Step 1: the model's top-level structure.
# --------------------------------------------------------------------------
# A top-level definition starts at column 0 with one of Sail's declaration
# keywords. Order matters: `function clause execute` must be tried before
# `function clause`, and that before `function`, or the longer forms are
# swallowed by the shorter prefix.
_VIS = r"(?:private\s+|public\s+)?"

_DEFS = [
    ("execute",       re.compile(r"^function clause execute\s+(\w+)")),
    ("function_clause", re.compile(r"^function clause\s+(\w+)")),
    # `private` and `public` are modifiers, not part of the keyword. Missing
    # them made the previous definition swallow the next one, which the
    # ownership cross-check caught as `lift_sstatus` being attributed to
    # `lower_mstatus`. The same omission cost the inventory script its PMP
    # address-match types.
    ("function",      re.compile(r"^" + _VIS + r"function\s+operator\s+(\S+)")),
    ("function",      re.compile(r"^" + _VIS + r"function\s+(\w+)")),
    ("mapping_clause", re.compile(r"^" + _VIS + r"mapping clause\s+(\w+)")),
    ("mapping",       re.compile(r"^" + _VIS + r"mapping\s+(\w+)")),
    ("union_clause",  re.compile(r"^union clause\s+(\w+)")),
    ("enum_clause",   re.compile(r"^enum clause\s+(\w+)")),
    ("val",           re.compile(r"^" + _VIS + r"val\s+(\w+)")),
    ("register",      re.compile(r"^register\s+(\w+)")),
    ("let",           re.compile(r"^" + _VIS + r"let\s+\(?\s*(\w+)")),
    ("type",          re.compile(r"^type\s+(\w+)")),
    ("struct",        re.compile(r"^" + _VIS + r"struct\s+(\w+)")),
    ("enum",          re.compile(r"^" + _VIS + r"enum\s+(\w+)")),
    ("overload",      re.compile(r"^overload\s+(\w+)")),
    ("scattered",     re.compile(r"^scattered\s+\w+\s+(\w+)")),
]

# `match <scrutinee> {` opening a block whose body is a list of arms. Only the
# form where the brace ends the line is recognised, which is how the model is
# written throughout; a single-line `match x { A => 1, B => 2 }` yields no arms
# and its spans fall back to owner-level attribution rather than being
# misattributed to whichever arm happened to match first.
_MATCH_OPEN = re.compile(r"\bmatch\s+([^{]*?)\s*\{\s*$")

# Any construct that makes the code below it conditional. Used to measure how
# many conditions must hold together for a span to execute -- see
# _index_conditions.
_COND_OPEN = re.compile(r"\bmatch\s+[^{]*\{\s*$|\bif\b.*\bthen\s*\{\s*$|"
                        r"\belse\s*\{\s*$|\bif\b[^;]*\{\s*$")


def _clean_cond(code):
    """The condition itself, without the syntax that opened its block."""
    t = code.strip()
    t = re.sub(r"^\}\s*", "", t)
    t = re.sub(r"\s*\bthen\s*\{\s*$", "", t)
    t = re.sub(r"\s*\{\s*$", "", t)
    return t.strip()[:80] or "else"

# An arm pattern: a bare constructor, a constructor with an argument list, a
# wildcard, or an or-pattern. Only the leading name is captured -- that is what
# the mnemonic map is keyed on.
_ARM = re.compile(r"^\s*([A-Za-z_]\w*|_)\s*(?:\([^)]*\))?\s*(?:\|[^=]*)?=>")

# The model's own authority on what an enum constructor is called in assembly.
_STR_MAP_OPEN = re.compile(r"^mapping\s+\w+\s*:\s*[^<]*<->\s*string\b")
_STR_MAP_ENTRY = re.compile(r'^\s*([A-Z]\w*)\s*<->\s*"([^"]+)"')


def _feature(relpath):
    """The reviewer-facing grouping for a model file.

    Taken from the model's directory layout, which is how the model itself
    organises extensions -- not from a table here that would drift from it.
    """
    parts = relpath.split("/")
    if parts[0] == "extensions" and len(parts) > 2:
        return parts[1]
    if parts[0] in ("sys", "pmp", "core", "vmem"):
        return parts[0]
    return parts[0]


class SailIndex:
    """Structural index of the Sail sources: definitions, match arms, names."""

    def __init__(self, model_dir):
        self.model_dir = model_dir
        self.defs = collections.defaultdict(list)   # file -> [(first,last,kind,name)]
        self.arms = collections.defaultdict(list)   # file -> [(first,last,scrutinee,pattern,depth)]
        self.lines = {}                             # file -> [str]
        self.conditions = {}                        # file -> {line: [guard, ...]}
        self.mnemonic = {}                          # (file, CTOR) -> "text"
        self.mnemonic_any = {}                      # CTOR -> "text"
        self.calls = collections.defaultdict(set)    # caller name -> callee names
        self.callers = collections.defaultdict(set)  # callee name -> caller names
        self.execute_of = {}                         # function name -> execute ctor
        self.exec_targets = collections.defaultdict(list)  # ctor -> [mnemonic]
        self.exec_params = collections.defaultdict(set)   # ctor -> {parameter names}
        self.exec_file = {}                              # ctor -> defining file
        for root, _d, files in os.walk(model_dir):
            for fn in sorted(files):
                if fn.endswith(".sail"):
                    p = os.path.join(root, fn)
                    self._index_file(os.path.relpath(p, model_dir), p)
        self._index_callgraph()

    # -- indexing ----------------------------------------------------------
    def _index_file(self, rel, path):
        lines = open(path, errors="ignore").read().splitlines()
        self.lines[rel] = lines

        starts = []
        for i, raw in enumerate(lines):
            if not raw[:1].strip():
                continue                             # not at column 0
            for kind, rx in _DEFS:
                m = rx.match(raw)
                if m:
                    starts.append((i, kind, m.group(1)))
                    if kind == "execute":
                        # The clause's own operand names. An arm only denotes an
                        # instruction when the thing being matched on is one of
                        # them -- see _is_opcode_match.
                        args = re.search(r"\(([^)]*)\)", raw)
                        if args:
                            self.exec_params[m.group(1)] |= {
                                a.strip() for a in args.group(1).split(",") if a.strip()}
                    break
        for n, (i, kind, name) in enumerate(starts):
            last = (starts[n + 1][0] - 1) if n + 1 < len(starts) else len(lines) - 1
            self.defs[rel].append((i + 1, last + 1, kind, name))

        self._index_conditions(rel, lines)
        self._index_arms(rel, lines)
        self._index_mnemonics(rel, lines)

    def _index_conditions(self, rel, lines):
        """The full chain of conditions guarding each line, outermost first.

        The nearest enclosing guard is enough to classify a span's stimulus, but
        not to generate for it. A span two levels deep needs condition A *and*
        condition B to hold at once, and a directive naming only B sends the
        generator after a state it will never reach -- which is the same thing a
        SystemVerilog `cross` expresses, arrived at from the model's control
        structure rather than from a coverpoint declaration.

        7.5% of the model's spans sit at depth 2 or more, concentrated in
        `pmaCheck`, `check_PTE_permission`, `pt_walk` and `legalize_satp` --
        precisely the privileged logic where the architecture branches on
        combinations rather than on single facts.
        """
        chains, depth, stack = {}, 0, []
        for i, raw in enumerate(lines, 1):
            code = raw.split("//")[0]
            # Closing braces that precede any opening brace on this line are
            # applied first. This is what separates `} else if c then {` -- an
            # alternative, which pops the previous condition and pushes its own
            # at the same depth -- from a genuine nesting, which only pushes.
            # Counting an else-if chain as nesting put `clint_store` at depth 8
            # and claimed eight conditions had to hold at once, when in truth
            # exactly one of eight alternatives does.
            head = code.split("{", 1)[0] if "{" in code else code
            depth -= head.count("}")
            while stack and depth < stack[-1][0]:
                stack.pop()
            chains[i] = [t for _d, t in stack]
            opened = _COND_OPEN.search(code)
            # The leading `}` were applied above; the rest of the line still is
            # not, so add back what was already taken off.
            depth += code.count("{") - code.count("}") + head.count("}")
            if opened:
                stack.append((depth, _clean_cond(code)))
        self.conditions[rel] = chains

    def _index_arms(self, rel, lines):
        """Recover `match` arms by brace depth.

        Depth is counted over comment-stripped lines. A `match ... {` that ends
        a line opens an arm scope whose arms sit at the depth immediately
        inside it; the scope closes when depth falls back. Arms are recorded
        innermost-last so a lookup can take the last enclosing one.
        """
        depth = 0
        stack = []          # [(body_depth, scrutinee, open_arm_or_None)]
        out = []
        for i, raw in enumerate(lines):
            code = raw.split("//")[0]
            # Close any match scopes this line has fallen out of, ending their
            # open arm at the previous line.
            while stack and depth < stack[-1][0]:
                body_depth, scrut, arm = stack.pop()
                if arm is not None:
                    out.append((arm[0], i, scrut, arm[1], body_depth))
            # An arm boundary: only at the exact depth of the innermost match
            # body, so a nested `if { ... }` inside an arm cannot start one.
            if stack and depth == stack[-1][0]:
                m = _ARM.match(code)
                if m:
                    body_depth, scrut, arm = stack[-1]
                    if arm is not None:
                        out.append((arm[0], i, scrut, arm[1], body_depth))
                    stack[-1] = (body_depth, scrut, (i + 1, m.group(1)))
            opened = _MATCH_OPEN.search(code)
            depth += code.count("{") - code.count("}")
            if opened:
                stack.append((depth, opened.group(1).strip(), None))
        while stack:
            body_depth, scrut, arm = stack.pop()
            if arm is not None:
                out.append((arm[0], len(lines), scrut, arm[1], body_depth))
        # Sorted by (start, depth) so the innermost enclosing arm is the last
        # match when scanning for containment.
        self.arms[rel] = sorted(out, key=lambda a: (a[0], a[4]))

    def _index_mnemonics(self, rel, lines):
        inside = False
        depth = 0
        for raw in lines:
            code = raw.split("//")[0]
            if not inside:
                if _STR_MAP_OPEN.match(code):
                    inside = True
                    depth = code.count("{") - code.count("}")
                continue
            m = _STR_MAP_ENTRY.match(code)
            if m:
                ctor, text = m.group(1), m.group(2)
                self.mnemonic[(rel, ctor)] = text
                self.mnemonic_any.setdefault(ctor, text)
            depth += code.count("{") - code.count("}")
            if depth <= 0:
                inside = False

    def _index_callgraph(self):
        """Who calls whom, and which instructions each helper sits under.

        A helper like `carryless_mul` carries no guard of its own -- it is
        uncovered because no instruction that calls it has been generated, and
        the plan is useless unless it can say *which* instruction that is. The
        call graph is built by scanning each definition's body for applications
        of names that are themselves defined in the model, then walked
        backwards to the `execute` clauses that reach it.

        Deliberately syntactic: it will over-approximate (a name mentioned but
        not called) rather than miss a caller, because a plan item that names
        one instruction too many costs a reader a moment and one too few sends
        them down a dead end.
        """
        known = {name for defs in self.defs.values()
                 for _f, _l, kind, name in defs
                 if kind in ("function", "function_clause", "val", "mapping",
                             "mapping_clause", "execute")}
        for rel, defs in self.defs.items():
            lines = self.lines[rel]
            for first, last, kind, name in defs:
                if kind not in ("function", "function_clause", "execute"):
                    continue
                body = "\n".join(l.split("//")[0]
                                 for l in lines[first:last])
                callees = {m for m in re.findall(r"\b([a-zA-Z_]\w*)\s*\(", body)
                           if m in known and m != name}
                self.calls[name] |= callees
                for c in callees:
                    self.callers[c].add(name)
                if kind == "execute":
                    self.execute_of[name] = name
                    self.exec_file.setdefault(name, rel)

        # Every mnemonic an execute clause can produce, from its own arms. This
        # is what makes the backwards walk land on something a person can type.
        for rel, arms in self.arms.items():
            for first, _last, scrut, pattern, _d in arms:
                kind, name = self.owner(rel, first)
                if kind != "execute" or pattern == "_":
                    continue
                if not self.is_opcode_match(name, scrut):
                    continue
                mnem = self.mnemonic.get((rel, pattern)) or self.mnemonic_any.get(pattern)
                if mnem and mnem not in self.exec_targets[name]:
                    self.exec_targets[name].append(mnem)

    def is_opcode_match(self, ctor, scrutinee):
        """Is this `match` discriminating the instruction, or something else?

        `function clause execute RTYPE(rs2, rs1, rd, op)` matching on `op`
        discriminates the instruction: each arm is `add`, `sub`, `sra`. The same
        clause matching on `plat_misaligned_access.amo` does not -- that is a
        platform setting, and its arms (`AccessFault`, `AlignmentException`) are
        not instructions however much they look like enum constructors with
        string mappings, because `platform_config.sail` gives those names to
        parse a configuration file.

        Requiring the scrutinee to be one of the clause's own parameters
        separates the two without a list of names to maintain.
        """
        if not scrutinee:
            return False
        return scrutinee.strip() in self.exec_params.get(ctor, ())

    def reaching_instructions(self, func, limit=8, depth=6):
        """Instructions whose execution can reach `func`, walking callers back.

        Returns (names, truncated, files) -- `files` being where those
        instructions are defined, which is what the sweep takes as an argument.
        Breadth-first so the nearest instructions
        are reported first, and capped because "reached by 300 instructions" is
        not a plan item -- it is a helper everything uses, and the cap is the
        honest way to say so.
        """
        seen, frontier, found, files = {func}, {func}, [], []
        for _ in range(depth):
            nxt = set()
            for f in frontier:
                for c in sorted(self.callers.get(f, ())):
                    if c in seen:
                        continue
                    seen.add(c)
                    if c in self.execute_of:
                        f = self.exec_file.get(c)
                        if f and f not in files:
                            files.append(f)
                        for m in (self.exec_targets.get(c) or [c]):
                            if m not in found:
                                found.append(m)
                    else:
                        nxt.add(c)
            if len(found) >= limit or not nxt:
                break
            frontier = nxt
        return found[:limit], len(found) > limit, files[:3]

    # -- lookup ------------------------------------------------------------
    def owner(self, rel, line):
        for first, last, kind, name in self.defs.get(rel, ()):
            if first <= line <= last:
                return kind, name
        return "", ""

    def arm(self, rel, line):
        """The innermost match arm containing `line`, or (None, None)."""
        best = None
        for first, last, scrut, pattern, depth in self.arms.get(rel, ()):
            if first <= line <= last and (best is None or depth >= best[4]):
                best = (first, last, scrut, pattern, depth)
        return (best[2], best[3]) if best else (None, None)

    def chain(self, rel, line):
        """[outermost condition, ..., innermost] guarding this line."""
        return self.conditions.get(rel, {}).get(line, [])

    def text(self, rel, line):
        ls = self.lines.get(rel, ())
        return ls[line - 1] if 0 < line <= len(ls) else ""

    def guard(self, rel, line, back=14):
        """Span line plus the nearest conditional headers above it."""
        ls = self.lines.get(rel, ())
        own = self.text(rel, line)
        start = max(0, line - 1 - back)
        above = [l for l in ls[start:line - 1] if re.search(r"\b(if|match|when)\b", l)]
        return own, "\n".join(above[-3:] + [own])


# --------------------------------------------------------------------------
# Step 4: what stimulus a non-instruction span requires.
# --------------------------------------------------------------------------
# Ordered, first match wins, most specific first. `directive` says what a
# generator has to arrange; `flags` gives the concrete invocation where one
# exists. Kept aligned with coverage_analysis.py, which classifies the same
# guards for the uncovered-only queue.
RULES = [
    ("naming_mapping",
     re.compile(r'^mapping clause \w+\s*=\s*\w+\s*<->\s*"'),
     "out of reach: renders an enum as text; execution never consults it "
     "(same argument as derive_span_exclusions.py, different declaration form)",
     None),

    ("model_callback",
     re.compile(r"^function\s+\w*_?callback\b|^function\s+\w+_callback"),
     "instrumentation hook, not architecture: reached whenever the event it "
     "observes occurs", None),

    ("hypervisor",
     re.compile(r"VirtualUser|VirtualSupervisor|Hypervisor|\bhgatp\b|\bhstatus\b"),
     "out of reach: Hypervisor is not implemented in this model", None),

    ("privilege",
     re.compile(r"^\s*(User|Supervisor|Machine)\s*=>"),
     "execute the same code at this privilege level",
     {"User": "opcode_sweep.py --isla-arg=--preload-xpp=mpp=u --isla-arg=--pmp-allow-all",
      "Supervisor": "opcode_sweep.py --isla-arg=--run-in-supervisor --isla-arg=--pmp-allow-all",
      "Machine": "(default: the sweep already runs in M-mode)"}),

    ("csr_field",
     re.compile(r"\b([mshv]?(?:status|envcfg|counteren|ideleg|edeleg|seccfg|"
                r"tvec|cause|epc|tval|ie|ip|atp|scratch))\s*\[\s*(\w+)\s*\]"),
     "preload this CSR field before the instruction under test", None),

    ("configuration",
     re.compile(r"\bconfig\s+([\w.]+)"),
     "select or build a model configuration where this key differs", None),

    ("extension_gate",
     re.compile(r"currentlyEnabled\s*\(\s*(Ext_\w+)\s*\)|hartSupports\s*\(\s*(Ext_\w+)"),
     "run with this extension enabled, and again with it disabled", None),

    ("xlen",
     re.compile(r"\bxlen\s*==\s*(\d+)|sizeof\(xlen\)|\bin32BitMode\b"),
     "run this target at the other XLEN", None),

    ("translation",
     re.compile(r"\bSv(32|39|48|57)\b|\bBare\b|satp\.\w+|\bPTE_|pte_is_"),
     "install a page table in this translation mode and fault or walk through it",
     None),

    ("pmp",
     re.compile(r"\bpmp\w*|\bTOR\b|\bNAPOT\b|\bNA4\b|\bOFF\b"),
     "program a PMP entry of this match type and access across its boundary",
     None),

    ("memory_op",
     re.compile(r"(Load|Store|Atomic|LoadReserved|StoreConditional|"
                r"InstructionFetch|CacheAccess)\s*\("),
     "issue this kind of memory access", None),

    ("fault_path",
     re.compile(r"Illegal_Instruction|internal_error|access_fault|Addr_Align|"
                r"_Page_Fault|accessFault|\bfailure\b|\bE_[A-Z]\w*"),
     "construct the faulting condition and assert the resulting trap", None),

    ("operand_value",
     re.compile(r"(!=|==)\s*zreg|\brs[12]\s*(!=|==)|\brd\s*(!=|==)|"
                r"\bzeros\(\)|\bshamt\b"),
     "vary the operand: the guard discriminates on a specific register or value",
     None),
]


# Sail names its string builders consistently, which is a more reliable signal
# than anything in the body -- `exceptionType_to_str` mentions every exception
# name in the architecture and would otherwise read as the richest fault-path
# item in the plan.
_DIAGNOSTIC_NAME = re.compile(r"_to_str$|_to_string$|^string_of_|_name_str$")


def classify(own, context, strict=False):
    """Classify the stimulus a span needs.

    `strict` matches only the span's own line. The two-pass structure matters:
    `currentlyEnabled(Ext_V)` appears in the context window of most of the
    vector model, and letting it match at 14 lines' distance labelled 2525
    spans "enable this extension" when the useful thing to say about most of
    them was which instruction calls them. Precise-on-the-line first, loose
    afterwards.
    """
    for name, rx, directive, flags in RULES:
        m = rx.search(own) if strict else (rx.search(own) or rx.search(context))
        if not m:
            continue
        detail = ""
        if name == "privilege":
            mode = (re.match(r"^\s*(\w+)", own.strip()) or [None, ""])[1] \
                if re.match(r"^\s*(\w+)", own.strip()) else ""
            detail = (flags or {}).get(mode, "")
        elif m.groups():
            detail = next((g for g in m.groups() if g), "")
        return name, directive, detail
    return "unclassified", "read the guard and decide what arranges it", ""


# --------------------------------------------------------------------------
# Assembling the plan.
# --------------------------------------------------------------------------
# branch_info's F rows carry the function name the span belongs to. That is an
# independent statement of ownership from the compiler itself, which makes it
# the right thing to check the containment index against.
_F_NAMED = re.compile(r'^F\s+\d+,\s*"([^"]*)",\s*"?([^",]*)"?,\s*(\d+),')


def validate_ownership(branch_info, index, include_library=False):
    """Check containment attribution against the names the compiler emitted.

    Containment has to carry the other 39466 spans, which have no name, so a
    silent parsing bug in the definition index would misattribute most of the
    plan with nothing to reveal it. The 1941 named spans are the control: if
    ownership agrees with them, the same index is trustworthy for the rest.

    Sail mangles some names (an `execute` clause compiles to a function named
    after its constructor), so agreement is checked with the mangling allowed
    for rather than by string equality.
    """
    agree = disagree = skipped = 0
    examples = []
    for line in open(branch_info):
        m = _F_NAMED.match(line)
        if not m:
            continue
        name, fname, lineno = m.group(1), m.group(2), int(m.group(3))
        # `F n, "name", "", 0, 0, 0, 0` has no source location; parse_spans
        # drops these from the span set, so they are not part of the plan and
        # must not be counted against the index either.
        if not fname or (not include_library and fname.startswith("/")):
            continue
        if fname not in index.defs:
            skipped += 1
            continue
        _kind, owner = index.owner(fname, lineno)
        if not owner:
            skipped += 1
        elif (owner == name or name.startswith(owner) or owner.startswith(name)
              or name.strip("()") == "operator " + owner
              or name.strip("()").replace("operator ", "") == owner):
            agree += 1
        else:
            disagree += 1
            if len(examples) < 5:
                examples.append(f"{fname}:{lineno} compiler says {name!r}, "
                                f"index says {owner!r}")
    return agree, disagree, skipped, examples


def check_freshness(branch_info, model_dir, coverage):
    """Refuse to describe a model that is not the one on disk.

    Everything here is derived, which makes the plan self-updating *provided the
    inputs are current*. They can silently not be:

      - `branch_info` is emitted by the coverage build. Edit the model without
        rebuilding and the target set is the previous model's -- new code is
        invisible, and the plan reports full coverage of behaviour that no
        longer exists. Nothing about the output would look wrong.
      - `sail_coverage` is a log of executed spans, matched by line number. Runs
        collected before a model edit line up against shifted spans, so status
        is joined to the wrong rows -- again silently, and in whichever
        direction the edit happened to shift things.

    Both failures produce a plausible plan, which is the dangerous kind. So the
    check is on by default and the run stops; `--allow-stale` is there for
    deliberately inspecting an older plan.
    """
    problems = []
    newest, newest_f = 0.0, None
    for root, _d, files in os.walk(model_dir):
        for fn in files:
            if fn.endswith(".sail"):
                t = os.path.getmtime(os.path.join(root, fn))
                if t > newest:
                    newest, newest_f = t, os.path.join(root, fn)
    bi = os.path.getmtime(branch_info)
    if newest > bi:
        problems.append(
            f"the model is newer than the span manifest:\n"
            f"    {os.path.relpath(newest_f, model_dir)} changed after "
            f"{os.path.basename(branch_info)} was built.\n"
            f"    The plan would describe the previous model, and new code "
            f"would not appear as a hole.\n"
            f"    Fix: cmake --build {os.path.dirname(branch_info)}")
    if coverage and os.path.exists(coverage):
        cov = os.path.getmtime(coverage)
        if cov < bi:
            problems.append(
                f"the coverage log predates the span manifest:\n"
                f"    {os.path.basename(coverage)} was written before "
                f"{os.path.basename(branch_info)} was built.\n"
                f"    Status would be joined to spans that have since moved.\n"
                f"    Fix: re-run coverage_report.py to replay the corpus.")
    return problems


def diff_plans(items, baseline_path):
    """What changed since a previous plan -- which is where new holes show up.

    A plan that only ever states the current position cannot answer "is this
    gap new?", and that is the question a model update raises. Items are keyed
    on (feature, target, stimulus) rather than on `item_id`, which is positional
    and shifts whenever anything is inserted above it.

    Four kinds of change are reported, and they mean different things:

      new hole      a target that did not exist before and is not covered --
                    almost always model code that was just added
      regression    a target that was fully covered and no longer is; the
                    corpus stopped reaching something it used to
      widened       a target that still exists but grew uncovered spans
      closed        work that was finished since the baseline
    """
    with open(baseline_path) as f:
        base = {(i["feature"], i["target"], i["stimulus"]): i
                for i in json.load(f)["items"]}
    cur = {(i["feature"], i["target"], i["stimulus"]): i for i in items}
    out = {"new hole": [], "regression": [], "widened": [], "closed": [],
           "new covered": [], "gone": []}
    for k, i in cur.items():
        if i["state"] == "unreachable":
            continue
        b = base.get(k)
        if b is None:
            out["new hole" if i["todo"] else "new covered"].append(i)
        elif b["todo"] == 0 and i["todo"] > 0:
            out["regression"].append(i)
        elif i["todo"] > b["todo"]:
            out["widened"].append(dict(i, was=b["todo"]))
        elif i["todo"] < b["todo"]:
            out["closed"].append(dict(i, was=b["todo"]))
    for k, b in base.items():
        if k not in cur and b.get("state") != "unreachable":
            out["gone"].append(b)
    return out


def load_exclusions(path):
    ranges = collections.defaultdict(list)
    if not path:
        return ranges
    for line in open(path):
        entry = line.split("#", 1)[0].strip()
        if not entry:
            continue
        fname, _, span = entry.rpartition(":")
        lo, _, hi = span.partition("-")
        try:
            ranges[fname].append((int(lo), int(hi or lo)))
        except ValueError:
            continue
    return ranges


def load_scope(path):
    if not path:
        return None, None
    entries = [p.strip() for line in open(path)
               for p in line.split("#", 1)[0].split(",") if p.strip()]
    return ([p for p in entries if not p.startswith("-")],
            [p[1:] for p in entries if p.startswith("-")])


def build_rows(index, spans, covered, excl, sweep_file):
    """One row per in-scope span, reasoned and given a status."""
    rows = []
    for kind, fname, l1, c1, l2, c2 in sorted(spans, key=lambda s: (s[1], s[2], s[3])):
        excluded = any(lo <= l1 <= hi for lo, hi in excl.get(fname, ()))
        own, ctx = index.guard(fname, l1)
        chain = index.chain(fname, l1)
        okind, oname = index.owner(fname, l1)
        scrut, pattern = index.arm(fname, l1)

        # An arm inside an execute clause names a concrete instruction when the
        # model's own mnemonic table knows its constructor. Same-file first: two
        # extensions can reuse a constructor name, and the local table is the
        # one that governs this clause.
        mnem = None
        if pattern and pattern != "_" and index.is_opcode_match(oname, scrut):
            mnem = index.mnemonic.get((fname, pattern)) or index.mnemonic_any.get(pattern)

        # The ladder, most specific first. Each rung is only taken when it can
        # name something concrete; anything that cannot falls through rather
        # than dressing a guess up as a directive.
        target = oname or fname
        generate_cmd = ""
        if okind == "execute" and mnem:
            # 1. A match arm inside an execute clause *is* one instruction.
            target = mnem
            stim, directive, detail = ("instruction_variant",
                                       "generate a test that executes this instruction",
                                       f"{oname} arm {pattern}")
        elif _DIAGNOSTIC_NAME.search(oname or ""):
            # 2. A string builder, identified by the model's own naming
            #    convention rather than by its text: the spans of interest are
            #    inside the body, where nothing says "this renders a value".
            stim = "diagnostic_string"
            directive = ("out of reach in normal execution: renders a value as "
                         "text for tracing, not architectural behaviour")
            detail = f"{oname}()"
        else:
            # 3. A discriminating guard on the span's own line.
            stim, directive, detail = classify(own, ctx, strict=True)
            if stim == "unclassified" and okind == "execute":
                # 4. Inside an execute clause with no arm: the base path.
                stim, directive = ("instruction",
                                   "generate any legal form of this instruction")
                detail = f"{oname} body" + (f", arm {pattern}" if pattern else "")
            elif stim == "unclassified" and okind in ("function", "function_clause"):
                # 5. A helper with no guard of its own is uncovered because
                #    nothing that calls it has been generated. Name the callers.
                insns, more, caller_files = index.reaching_instructions(oname)
                if insns:
                    stim = "helper_function"
                    directive = ("execute an instruction that calls this helper: "
                                 + ", ".join(insns) + (", ..." if more else ""))
                    detail = f"{oname}() reached via {len(insns)}"
                    if more:
                        detail += "+"
                    detail += " instruction(s)"
                    if caller_files:
                        # Override the file-level default: sweeping the helper's
                        # own file generates nothing, because a helper has no
                        # encoding. The instruction that calls it does.
                        generate_cmd = "opcode_sweep.py " + " ".join(caller_files)
                else:
                    # 6. Loose match, then honest failure.
                    stim, directive, detail = classify(own, ctx)
            elif stim == "unclassified":
                stim, directive, detail = classify(own, ctx)

        rows.append(collections.OrderedDict(
            feature=_feature(fname),
            target=target,
            owner=oname,
            owner_kind=okind,
            stimulus=stim,
            directive=directive,
            detail=detail,
            status="excluded" if excluded
                   else ("covered" if (kind, fname, l1, c1, l2, c2) in covered
                         else "todo"),
            span_kind=kind,
            file=fname,
            line=l1,
            match_on=scrut or "",
            arm=pattern or "",
            nesting=len(chain),
            # The conjunction a generator has to satisfy, outermost first. One
            # row of the plan, not several: a nested span is a single target
            # that happens to need several things true at once.
            conditions=" AND ".join(chain)[:400],
            source=own.strip()[:140],
            generate=generate_cmd or sweep_file.get(fname, ""),
        ))
    return rows


# Stimuli whose directive is "this cannot be reached by executing a program".
# Distinct from the derive_span_exclusions.py list in one important way: those
# spans leave the denominator, these stay in it. The exclusions file reasons
# structurally about a whole declaration form; these are single spans reasoned
# about individually, and a per-span judgement is not strong enough to remove
# anything from a coverage figure. So they are reported, and not scheduled.
OUT_OF_REACH = {"naming_mapping", "hypervisor", "diagnostic_string"}


def build_items(rows):
    """Roll spans up into the unit a person schedules.

    A plan item is one (feature, target, stimulus) group. `sra` is one item of
    four spans, not four items; `pmpCheck` under a TOR guard is one item
    regardless of how many expressions the arm compiles to. Excluded spans are
    counted but kept out of both numerator and denominator, so an item that is
    entirely unreachable reports 0/0 rather than a false shortfall.
    """
    groups = collections.OrderedDict()
    for r in rows:
        key = (r["feature"], r["target"], r["stimulus"])
        g = groups.setdefault(key, {
            "feature": r["feature"], "target": r["target"],
            "stimulus": r["stimulus"], "directive": r["directive"],
            "detail": r["detail"], "file": r["file"], "first_line": r["line"],
            "generate": r["generate"],
            "spans": 0, "covered": 0, "excluded": 0, "todo": 0,
        })
        g["spans"] += 1
        if r["nesting"] > g.get("max_nesting", 0):
            g["max_nesting"] = r["nesting"]
            g["conditions"] = r["conditions"]
        g[r["status"]] = g.get(r["status"], 0) + 1
        g["first_line"] = min(g["first_line"], r["line"])
    items = []
    for i, g in enumerate(groups.values(), 1):
        live = g["spans"] - g["excluded"]
        g["item_id"] = "TP%04d" % i
        g["reachable_spans"] = live
        g["percent"] = round(100.0 * g["covered"] / live, 1) if live else 0.0
        g["state"] = ("unreachable" if live == 0 else
                      "complete" if g["todo"] == 0 else
                      "out of reach" if g["stimulus"] in OUT_OF_REACH else
                      "partial" if g["covered"] else "not started")
        items.append(collections.OrderedDict(
            item_id=g["item_id"], feature=g["feature"], target=g["target"],
            stimulus=g["stimulus"], state=g["state"],
            spans=g["reachable_spans"], covered=g["covered"], todo=g["todo"],
            percent=g["percent"], excluded=g["excluded"],
            max_nesting=g.get("max_nesting", 0),
            cross="yes" if g.get("max_nesting", 0) >= 2 else "",
            conditions=g.get("conditions", ""),
            directive=g["directive"], detail=g["detail"],
            source=f"{g['file']}:{g['first_line']}", generate=g["generate"],
        ))
    # Worst-covered largest items first: that ordering is the schedule.
    items.sort(key=lambda it: (it["state"] == "complete",
                               it["state"] in ("unreachable", "out of reach"),
                               it["percent"], -it["spans"]))
    return items


def write_markdown(path, items, rows, meta, changes=None):
    live = [it for it in items if it["state"] != "unreachable"]
    tot = sum(it["spans"] for it in live)
    cov = sum(it["covered"] for it in live)
    by_feature = collections.OrderedDict()
    for it in live:
        f = by_feature.setdefault(it["feature"],
                                  {"items": 0, "done": 0, "spans": 0, "covered": 0})
        f["items"] += 1
        f["done"] += it["state"] == "complete"
        f["spans"] += it["spans"]
        f["covered"] += it["covered"]

    L = []
    L.append("# Testplan derived from the Sail model\n")
    L.append(f"Generated by `testplan.py` from `{meta['branch_info']}`"
             f" and `{meta['model_dir']}`.\n")
    L.append("Every row below was derived from the model sources. Nothing here is "
             "hand-listed, so\nthe plan cannot omit a behaviour the model "
             "implements: the target set *is* the set of\nspans the model "
             "compiles.\n")
    if meta["coverage"]:
        L.append(f"Status joined from `{meta['coverage']}`.\n")
    else:
        L.append("No coverage file supplied, so every item reads *not started*. "
                 "This is the plan\nas it stands before any test runs.\n")

    L.append("\n## Where the plan stands\n")
    L.append(f"| | |\n|---|---|")
    L.append(f"| Plan items | {len(live)} |")
    L.append(f"| Items complete | {sum(1 for it in live if it['state']=='complete')} |")
    L.append(f"| Items partial | {sum(1 for it in live if it['state']=='partial')} |")
    L.append(f"| Items not started | {sum(1 for it in live if it['state']=='not started')} |")
    L.append(f"| Items out of reach | {sum(1 for it in live if it['state']=='out of reach')} |")
    L.append(f"| Reachable spans | {tot} |")
    L.append(f"| Spans covered | {cov} ({100*cov/tot if tot else 0:.1f}%) |")
    L.append(f"| Items held unreachable | {len(items)-len(live)} |")

    L.append("\n## By feature\n")
    L.append("| Feature | Items | Complete | Spans | Covered | % |")
    L.append("|---|---:|---:|---:|---:|---:|")
    for f, d in sorted(by_feature.items(), key=lambda kv: -kv[1]["spans"]):
        L.append(f"| {f} | {d['items']} | {d['done']} | {d['spans']} | "
                 f"{d['covered']} | {100*d['covered']/d['spans'] if d['spans'] else 0:.1f}% |")

    L.append("\n## By stimulus\n")
    L.append("What a generator has to arrange, and how much rides on it.\n")
    L.append("| Stimulus | Items | Spans | Covered | % | What it takes |")
    L.append("|---|---:|---:|---:|---:|---|")
    by_stim = collections.OrderedDict()
    for it in live:
        s = by_stim.setdefault(it["stimulus"],
                               {"items": 0, "spans": 0, "covered": 0,
                                "directive": it["directive"]})
        s["items"] += 1
        s["spans"] += it["spans"]
        s["covered"] += it["covered"]
    for s, d in sorted(by_stim.items(), key=lambda kv: -kv[1]["spans"]):
        L.append(f"| `{s}` | {d['items']} | {d['spans']} | {d['covered']} | "
                 f"{100*d['covered']/d['spans'] if d['spans'] else 0:.1f}% | {d['directive']} |")

    if changes:
        new = changes["new hole"] + changes["regression"]
        L.append("\n## What changed since the baseline\n")
        if new:
            L.append(f"**{len(changes['new hole'])} new holes** and "
                     f"**{len(changes['regression'])} regressions**. A new hole is "
                     f"a target the model\nhas gained that nothing covers; a "
                     f"regression is one the corpus used to reach and no\nlonger "
                     f"does. Both are worth more attention than the standing "
                     f"backlog below.\n")
            L.append("| Kind | Feature | Target | Stimulus | To do | Source |")
            L.append("|---|---|---|---|---:|---|")
            for kind in ("new hole", "regression"):
                for it in sorted(changes[kind], key=lambda i: -i["todo"])[:25]:
                    L.append(f"| {kind} | {it['feature']} | `{it['target']}` | "
                             f"{it['stimulus']} | {it['todo']} | {it['source']} |")
        else:
            L.append("No new holes and no regressions against the baseline.\n")
        if changes["closed"]:
            L.append(f"\n{len(changes['closed'])} items advanced and "
                     f"{len(changes['gone'])} disappeared from the model.\n")

    crosses = [i for i in live if i["cross"] and i["todo"]]
    L.append("\n## Targets that need several conditions at once\n")
    L.append(f"{len(crosses)} outstanding items have at least one span guarded by "
             f"two or more nested\nconditions. These are the model's own crosses: "
             f"reaching the span needs A *and* B to\nhold together, so a "
             f"directive naming only the innermost condition sends the\n"
             f"generator after a state it will never reach.\n")
    L.append("| Item | Feature | Target | Depth | To do | Conditions that must hold together |")
    L.append("|---|---|---|---:|---:|---|")
    for it in sorted(crosses, key=lambda i: (-i["max_nesting"], -i["todo"]))[:20]:
        cond = it["conditions"].replace("|", "\\|")[:150]
        L.append(f"| {it['item_id']} | {it['feature']} | `{it['target']}` | "
                 f"{it['max_nesting']} | {it['todo']} | `{cond}` |")
    L.append("\nA note on what this is *not*: these are crosses in the model's "
             "control flow. ACT4's\n`cr_rs1_rs2_edges` crosses operand *values* "
             "-- 11 edge values on rs1 against 11 on\nrs2 -- and Sail has no "
             "branch on operand values for arithmetic, so `add` is one span\n"
             "that any single execution covers. That axis is invisible to "
             "source coverage by\nconstruction, and is why the RVVI/SVA "
             "coverpoints are measured alongside rather than\nreplaced by "
             "this plan.\n")

    L.append("\n## What to generate next\n")
    L.append("The 40 largest items with work outstanding, worst-covered first. "
             "This ordering is\nthe schedule: the top of this table is the "
             "cheapest coverage per unit of work.\n")
    L.append("| Item | Feature | Target | Stimulus | Spans | Done | To do | Source |")
    L.append("|---|---|---|---|---:|---:|---:|---|")
    for it in [i for i in live if i["todo"] and i["state"] != "out of reach"][:40]:
        L.append(f"| {it['item_id']} | {it['feature']} | `{it['target']}` | "
                 f"{it['stimulus']} | {it['spans']} | {it['covered']} | "
                 f"{it['todo']} | {it['source']} |")

    done = [i for i in live if i["state"] == "complete"]
    L.append(f"\n## Already covered\n\n{len(done)} items are fully covered by the "
             f"existing corpus and need no new\ntest. They stay in the plan: an "
             f"item that disappears when it passes cannot be\nshown to have "
             f"regressed.\n")

    oor = [i for i in live if i["state"] == "out of reach"]
    if oor:
        n_spans = sum(i["todo"] for i in oor)
        L.append(f"## Out of reach, but still counted\n\n{len(oor)} items "
                 f"({n_spans} spans) cannot be reached by executing a program: "
                 f"enum-to-text\nmappings, diagnostic string builders, and "
                 f"Hypervisor code the model does not implement.\n\nThey are "
                 f"left in the denominator above. `derive_span_exclusions.py` "
                 f"removes whole\ndeclaration forms it can reason about "
                 f"structurally; these are per-span judgements,\nwhich is not "
                 f"a strong enough basis to move a coverage figure. So they are "
                 f"reported\nhere and kept out of the schedule, not deducted.\n")
        L.append("| Feature | Target | Spans | Why | Source |")
        L.append("|---|---|---:|---|---|")
        for it in sorted(oor, key=lambda i: -i["todo"])[:15]:
            L.append(f"| {it['feature']} | `{it['target']}` | {it['todo']} | "
                     f"{it['stimulus']} | {it['source']} |")
        L.append("")

    unreach = [i for i in items if i["state"] == "unreachable"]
    if unreach:
        L.append(f"## Held unreachable\n\n{len(unreach)} items consist entirely of "
                 f"spans that instruction execution cannot\nreach, per "
                 f"`derive_span_exclusions.py`. They are listed rather than "
                 f"deleted so the\nexclusion stays auditable.\n")
        L.append("| Feature | Target | Spans | Source |")
        L.append("|---|---|---:|---|")
        for it in unreach[:20]:
            L.append(f"| {it['feature']} | `{it['target']}` | {it['excluded']} | {it['source']} |")
        if len(unreach) > 20:
            L.append(f"\n...and {len(unreach)-20} more, in the CSV.\n")

    open(path, "w").write("\n".join(L) + "\n")


def write_xlsx(path, items, rows):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        print("openpyxl not available; skipped .xlsx", file=sys.stderr)
        return
    wb = Workbook()

    def sheet(ws, data, widths=60):
        ws.append(list(data[0].keys()))
        for r in data:
            ws.append(list(r.values()))
        for c in ws[1]:
            c.fill = PatternFill("solid", fgColor="1B3A6B")
            c.font = Font(color="FFFFFF", bold=True)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for i, key in enumerate(data[0].keys(), 1):
            w = max([len(str(key))] + [len(str(r[key])) for r in data[:2000]]) + 2
            ws.column_dimensions[get_column_letter(i)].width = min(w, widths)

    ws = wb.active
    ws.title = "plan items"
    sheet(ws, items)
    if rows:
        sheet(wb.create_sheet("spans"), rows)
    wb.save(path)


def main():
    ap = argparse.ArgumentParser(
        description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--model-dir", default=os.path.join(paths.SAIL_RISCV, "model"))
    ap.add_argument("--branch-info", default=paths.BRANCH_INFO)
    ap.add_argument("--coverage", nargs="?", const=paths.COVERAGE_FILE,
                    help="a sail_coverage file to join in as status; without it "
                         "the plan is the target set alone")
    ap.add_argument("--exclude-spans", help="ranges from derive_span_exclusions.py")
    ap.add_argument("--scope", help="same scope file coverage_report.py takes")
    ap.add_argument("-o", "--out-prefix", default="testplan")
    ap.add_argument("--xlsx", action="store_true")
    ap.add_argument("--json-spans", action="store_true",
                    help="include the per-span rows in the JSON as well. Off by "
                         "default: they duplicate <prefix>-spans.csv and take the "
                         "file from ~1MB to ~10MB, which matters because this "
                         "JSON is the checked-in --baseline.")
    ap.add_argument("--html", action="store_true",
                    help="also write a self-contained HTML plan with a filterable "
                         "item table -- the form to hand a reviewer, as opposed to "
                         "the CSV, which is the form to work from.")
    ap.add_argument("--baseline", metavar="testplan.json",
                    help="a previous plan to diff against. Reports which gaps "
                         "are NEW -- targets the model has gained, and targets "
                         "the corpus has stopped reaching -- as opposed to the "
                         "standing backlog, which a bare plan cannot separate.")
    ap.add_argument("--fail-on-new-holes", action="store_true",
                    help="exit non-zero if the diff shows a new hole or a "
                         "regression, so a model bump fails CI rather than "
                         "quietly enlarging the backlog. Needs --baseline.")
    ap.add_argument("--allow-stale", action="store_true",
                    help="proceed even when the span manifest is older than the "
                         "model, or the coverage log older than the manifest. "
                         "For inspecting an old plan on purpose; the numbers are "
                         "not about the model currently on disk.")
    ap.add_argument("--include-library", action="store_true",
                    help="keep Sail standard-library spans; excluded by default "
                         "because they measure the library, not the model")
    args = ap.parse_args()

    for p, what in ((args.branch_info, "branch_info manifest"),
                    (args.model_dir, "model directory")):
        if not os.path.exists(p):
            sys.exit(f"missing {what}: {p}")

    stale = check_freshness(args.branch_info, args.model_dir, args.coverage)
    if stale:
        for pr in stale:
            print(f"STALE INPUT: {pr}", file=sys.stderr)
        if not args.allow_stale:
            sys.exit("\nRefusing to write a plan from stale inputs -- it would "
                     "look correct and\ndescribe a different model. Rebuild as "
                     "above, or pass --allow-stale if that is\nwhat you meant.")
        print("--allow-stale: continuing anyway\n", file=sys.stderr)

    spans = parse_spans(args.branch_info)
    if not args.include_library:
        spans = {s for s in spans if not s[1].startswith("/")}
    prefixes, excludes = load_scope(args.scope)
    if prefixes:
        spans = {s for s in spans
                 if any(s[1].startswith(p) for p in prefixes)
                 and not any(s[1].startswith(x) for x in excludes)}
    if not spans:
        sys.exit("no spans in scope -- check --scope and --branch-info")

    covered = parse_spans(args.coverage) if args.coverage else set()
    excl = load_exclusions(args.exclude_spans)

    print(f"indexing {args.model_dir} ...", flush=True)
    index = SailIndex(os.path.abspath(args.model_dir))
    print(f"  {sum(len(v) for v in index.defs.values())} top-level definitions, "
          f"{sum(len(v) for v in index.arms.values())} match arms, "
          f"{len(index.mnemonic)} mnemonic bindings")

    agree, disagree, skipped, examples = validate_ownership(
        args.branch_info, index, args.include_library)
    total_named = agree + disagree
    pct = 100.0 * agree / total_named if total_named else 0.0
    print(f"  ownership cross-check: {agree}/{total_named} named spans agree "
          f"({pct:.1f}%), {skipped} not comparable")
    for e in examples:
        print(f"    disagreement: {e}")
    if total_named and pct < 90:
        print("  WARNING: the definition index disagrees with the compiler on more\n"
              "  than one named span in ten. Attribution for the unnamed spans --\n"
              "  which is most of the plan -- should not be trusted until this is\n"
              "  understood.", file=sys.stderr)

    # The command that generates for a given model file, so a plan row carries
    # its own reproduction rather than referring to a runbook.
    sweep_file = {f: f"opcode_sweep.py {f}" for f in {s[1] for s in spans}}

    rows = build_rows(index, spans, covered, excl, sweep_file)
    items = build_items(rows)

    # Computed before anything is written: the markdown carries the diff, and
    # a baseline that fails to load should do so before files are replaced.
    changes = diff_plans(items, args.baseline) if args.baseline else None

    with open(args.out_prefix + "-spans.csv", "w", newline="") as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0].keys()))
        w.writeheader()
        w.writerows(rows)
    with open(args.out_prefix + ".csv", "w", newline="") as f:
        w = csv.DictWriter(f, fieldnames=list(items[0].keys()))
        w.writeheader()
        w.writerows(items)
    meta = {"branch_info": args.branch_info, "model_dir": args.model_dir,
            "coverage": args.coverage, "scope": args.scope,
            "exclusions": args.exclude_spans}
    with open(args.out_prefix + ".json", "w") as f:
        # Items only unless asked: --baseline reads nothing else, and the span
        # rows are already in <prefix>-spans.csv in a form that diffs.
        payload = {"provenance": meta, "items": items}
        if args.json_spans:
            payload["spans"] = rows
        json.dump(payload, f, indent=1)
    write_markdown(args.out_prefix + ".md", items, rows, meta, changes)
    if args.html:
        import testplan_html
        with open(args.out_prefix + ".html", "w") as f:
            f.write(testplan_html.render(items, rows, meta, changes))
    if args.xlsx:
        write_xlsx(args.out_prefix + ".xlsx", items, rows)

    if args.baseline:
        print(f"\n{'=' * 62}\nchanges since {args.baseline}\n{'=' * 62}")
        order = ["new hole", "regression", "widened", "closed", "new covered", "gone"]
        for kind in order:
            group = changes[kind]
            if not group:
                continue
            print(f"\n{kind.upper()}  ({len(group)})")
            for it in sorted(group, key=lambda i: -i.get("todo", 0))[:12]:
                was = f"  (was {it['was']} to do)" if "was" in it else ""
                print(f"  {it['feature']:<14} {it['target']:<26} "
                      f"{it['stimulus']:<20} {it['todo']:>4} to do{was}")
            if len(group) > 12:
                print(f"  ... and {len(group)-12} more")
        if not any(changes[k] for k in order):
            print("\nno change")

    live = [it for it in items if it["state"] != "unreachable"]
    tot = sum(it["spans"] for it in live)
    cov = sum(it["covered"] for it in live)
    print(f"\n{len(rows)} spans -> {len(items)} plan items "
          f"({len(live)} reachable, {len(items)-len(live)} held unreachable)")
    print(f"{args.out_prefix}.{{md,csv,json}} + {args.out_prefix}-spans.csv"
          + (f" + {args.out_prefix}.xlsx" if args.xlsx else "")
          + (f" + {args.out_prefix}.html" if args.html else ""))
    print(f"\n{'state':<14}{'items':>7}")
    for st, n in collections.Counter(it["state"] for it in items).most_common():
        print(f"{st:<14}{n:>7}")
    print(f"\nreachable spans {cov}/{tot} = {100*cov/tot if tot else 0:.1f}%")
    print(f"\n{'stimulus':<20}{'items':>7}{'spans':>8}{'todo':>8}")
    agg = collections.defaultdict(lambda: [0, 0, 0])
    for it in live:
        a = agg[it["stimulus"]]
        a[0] += 1
        a[1] += it["spans"]
        a[2] += it["todo"]
    for s, (n, sp, td) in sorted(agg.items(), key=lambda kv: -kv[1][2]):
        print(f"{s:<20}{n:>7}{sp:>8}{td:>8}")

    if args.fail_on_new_holes:
        if not changes:
            sys.exit("--fail-on-new-holes needs --baseline")
        n = len(changes["new hole"]) + len(changes["regression"])
        if n:
            sys.exit(f"\n{n} new hole(s)/regression(s) against the baseline.")


if __name__ == "__main__":
    main()
